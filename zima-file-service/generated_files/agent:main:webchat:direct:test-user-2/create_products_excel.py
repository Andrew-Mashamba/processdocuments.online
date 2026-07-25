#!/usr/bin/env python3

import sys
import os
sys.path.append('/Volumes/DATA/QWEN/zima-file-service')

# Add current directory to path for imports
current_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, current_dir)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"], check=True)
    import openpyxl
    from openpyxl.styles import Font, PatternFill

def create_products_excel():
    """Create an Excel file with 10 random product names and prices."""

    # Create a new workbook and select the active worksheet
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product List"

    # Define headers
    headers = ["Product Name", "Price"]

    # Define product data
    products = [
        ["Wireless Bluetooth Headphones", "$89.99"],
        ["Gaming Mechanical Keyboard", "$129.50"],
        ["Smart Fitness Watch", "$249.00"],
        ["USB-C Phone Charger", "$24.95"],
        ["Portable Bluetooth Speaker", "$65.75"],
        ["LED Desk Lamp", "$42.30"],
        ["Ergonomic Office Chair", "$399.99"],
        ["Coffee Machine", "$156.80"],
        ["External Hard Drive 1TB", "$78.45"],
        ["Noise-Cancelling Earbuds", "$179.95"]
    ]

    # Add headers to the worksheet
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    # Add product data
    for row_idx, product in enumerate(products, 2):
        for col_idx, value in enumerate(product, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Auto-fit column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save the file
    output_path = "/Volumes/DATA/QWEN/zima-file-service/generated_files/random_products.xlsx"
    wb.save(output_path)

    print(f"✅ Excel file created successfully: {output_path}")
    print(f"📊 Contains {len(products)} products with names and prices")

    return output_path

if __name__ == "__main__":
    create_products_excel()